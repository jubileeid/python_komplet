from openpyxl import Workbook, load_workbook
from openpyxl.chart import Reference, LineChart

wb = Workbook()
ws = wb.active
ws.title = "Chart"

ws.append(['Bulan','Susu','Sirup'])
ws.append(['Jan',100,200])
ws.append(['Feb',40,120])
ws.append(['Mar',80,60])
ws.append(['Apr',70,70])
ws.append(['Mei',100,120])
ws.append(['Jun',120,160])
ws.append(['Jul',110,80])
ws.append(['Agst',60,70])

values = Reference(ws, min_col=2, min_row=1, max_col=3, max_row=9)
x_values = Reference(ws, range_string="chart!A2:A9")

chart = LineChart()
chart.add_data(values, titles_from_data=True)

chart.set_categories(x_values)

chart.title = "Penjualan Susu dan Sirup"
chart.x_axis.title = "Bulan"
chart.y_axis.title = "Penjualan Susu dan Sirup"
chart.legend.position = "b"

ws.add_chart(chart, 'H1')

wb.save("chart.xlsx")
